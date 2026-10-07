# -*- coding: utf-8 -*-
"""
Builds IntuneRemediations-Poster.pptx - a single A1 landscape (33.1 x 23.4 in)
poster summarizing the pipeline. Export to PNG afterwards with PowerPoint:

    python Tests\\Build-Poster.py
    (then see Tests\\Export-PosterPng.ps1, or open the pptx and export manually)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "IntuneRemediations-Poster.pptx")

PRIMARY   = RGBColor(0x00, 0x78, 0xD4)
PRIMARY_D = RGBColor(0x00, 0x5A, 0x9E)
NAVY      = RGBColor(0x1B, 0x3A, 0x5B)
CARD_BG   = RGBColor(0xEE, 0xF3, 0xFA)
CARD_BD   = RGBColor(0xC9, 0xD6, 0xE8)
TEXT      = RGBColor(0x22, 0x2E, 0x3B)
MUTED     = RGBColor(0x5B, 0x6B, 0x7C)
GREEN     = RGBColor(0x10, 0x7C, 0x10)
GREEN_BG  = RGBColor(0xE7, 0xF3, 0xE7)
ORANGE    = RGBColor(0xC4, 0x3E, 0x1C)
ORANGE_BG = RGBColor(0xFB, 0xEF, 0xE8)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT     = RGBColor(0xB9, 0xD6, 0xF2)
GREY_ARR  = RGBColor(0x8F, 0xA8, 0xC2)
FONT      = "Calibri"

prs = Presentation()
prs.slide_width = Inches(33.11)
prs.slide_height = Inches(23.39)
s = prs.slides.add_slide(prs.slide_layouts[6])

W, H = 33.11, 23.39
M = 1.0                      # side margin
CW = W - 2 * M               # content width = 31.11


def _para(tf, first, text, size, color, bold=False, align=PP_ALIGN.LEFT,
          space_after=6, line_spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = color
    return p


def box(x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def text_at(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        size, color, bold = ln[1], ln[2], ln[3]
        _para(tf, i == 0, ln[0], size, color, bold, align,
              space_after=ln[4] if len(ln) > 4 else 6)
    return tb


def card(x, y, w, h, title, lines, tcolor=PRIMARY, fill=CARD_BG,
         border=CARD_BD, tsize=30, bsize=22):
    c = box(x, y, w, h, fill, border)
    tf = c.text_frame
    tf.margin_left = Inches(0.32)
    tf.margin_right = Inches(0.32)
    tf.margin_top = Inches(0.26)
    tf.margin_bottom = Inches(0.22)
    first = True
    if title:
        _para(tf, first, title, tsize, tcolor, True, PP_ALIGN.CENTER,
              space_after=12)
        first = False
    for ln in lines:
        lead, rest = (ln if isinstance(ln, tuple) else (None, ln))
        p = _para(tf, first, "", bsize, TEXT, False, PP_ALIGN.LEFT,
                  space_after=7, line_spacing=1.02)
        first = False
        if lead:
            r = p.add_run()
            r.text = lead
            r.font.bold = True
            r.font.size = Pt(bsize)
            r.font.name = FONT
            r.font.color.rgb = TEXT
        r = p.add_run()
        r.text = rest
        r.font.size = Pt(bsize)
        r.font.name = FONT
        r.font.color.rgb = TEXT
    return c


def arrow(x, y, w, h, shape=MSO_SHAPE.RIGHT_ARROW, color=GREY_ARR):
    return box(x, y, w, h, color, None, shape)


def table(x, y, col_widths, data, fsize=20, header_h=0.52, row_h=0.62,
          header_fill=PRIMARY):
    rows, cols = len(data), len(data[0])
    gf = s.shapes.add_table(rows, cols, Inches(x), Inches(y),
                            Inches(sum(col_widths)),
                            Inches(header_h + (rows - 1) * row_h))
    tbl = gf.table
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for ri in range(1, rows):
        tbl.rows[ri].height = Inches(row_h)
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.14)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = (
                    RGBColor(0xF5, 0xF8, 0xFC) if ri % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.size = Pt(fsize)
            r.font.bold = (ri == 0)
            r.font.name = FONT
            r.font.color.rgb = WHITE if ri == 0 else TEXT
    return tbl


# ============================================================ header
box(0, 0, W, 3.15, NAVY, None, MSO_SHAPE.RECTANGLE)
box(0, 3.15, W, 0.10, PRIMARY, None, MSO_SHAPE.RECTANGLE)
text_at(M, 0.55, 24.5, 1.3,
        [("Intune Patch & App Remediation Pipeline", 62, WHITE, True, 0)])
text_at(M, 1.85, 24.5, 1.1,
        [("Automated detection, error-code grouping and safe remediation "
          "of Windows Update and app installation failures", 27, LIGHT,
          False, 0)])
text_at(W - M - 8.5, 1.9, 8.5, 1.0,
        [("Repository: C:\\Users\\SAM\\Documents\\IntuneRemediations", 21,
          RGBColor(0x9F, 0xC4, 0xE8), False, 4),
         ("Technical deck (21 slides) + this poster  \u00b7  October 2026", 21,
          RGBColor(0x9F, 0xC4, 0xE8), False, 0)],
        align=PP_ALIGN.RIGHT)

# ============================================================ 1 - closed loop
text_at(M, 3.65, CW, 0.75,
        [("1  \u00b7  The closed loop", 36, PRIMARY_D, True, 0)])

flow_y, flow_h, bw = 4.35, 2.55, 5.30
gap = (CW - 5 * bw) / 4.0
boxes = [
    ("Azure Automation", [
        "PowerShell runbook, daily schedule",
        "Managed identity \u2013 no secrets stored",
    ], PRIMARY_D),
    ("Microsoft Graph", [
        "read failure states (paged, retried)",
        "create/update groups + members",
    ], PRIMARY_D),
    ("Entra error-code groups", [
        "Intune-PatchError-{code}",
        "Intune-AppError-{code}",
    ], PRIMARY_D),
    ("Intune Remediations", [
        "one assignment per group",
        "detection exit 0/1 \u2192 fix runs",
    ], GREEN),
    ("Windows devices", [
        "fix runs as SYSTEM via IME",
        "State|Detail lands in the report",
    ], PRIMARY_D),
]
for i, (t, lines, tc) in enumerate(boxes):
    x = M + i * (bw + gap)
    card(x, flow_y, bw, flow_h, t, lines, tcolor=tc, tsize=26, bsize=21)
    if i < 4:
        arrow(x + bw + 0.12, flow_y + (flow_h - 0.62) / 2, gap - 0.24, 0.62)

# return path: down at the right, strip flows back, up at the left
arrow(M + 4 * (bw + gap) + bw / 2 - 0.32, flow_y + flow_h + 0.07, 0.64, 0.50,
      MSO_SHAPE.DOWN_ARROW)
loop = box(M, flow_y + flow_h + 0.62, CW, 0.95, CARD_BG, CARD_BD)
tf = loop.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
_para(tf, True,
      "\u21ba  Closed loop: remediated devices change state \u2192 Graph reflects it "
      "\u2192 the next daily run empties the group \u2192 the assignment stops applying",
      25, NAVY, True, PP_ALIGN.CENTER, space_after=0)
arrow(M + bw / 2 - 0.32, flow_y + flow_h + 1.62, 0.64, 0.50,
      MSO_SHAPE.UP_ARROW)

# ============================================================ 2 - details
sec_y = 9.25
text_at(M, sec_y, CW, 0.75,
        [("2  \u00b7  How it works \u2013 sync, safety, contract", 36,
          PRIMARY_D, True, 0)])
col_y, col_h = sec_y + 0.8, 7.8
cgap = 0.7
cw = (CW - 2 * cgap) / 3.0

card(M, col_y, cw, col_h, "Daily group sync \u2013 9 steps", [
    ("1  ", "auth \u2013 managed identity, or -Interactive locally"),
    ("2  ", "read paged windowsUpdateStates / deviceStatuses"),
    ("3  ", "filter status = failed, bucket rows by errorCode"),
    ("4  ", "find or create Intune-PatchError-{code} groups"),
    ("5  ", "guard: groups without the automation marker are skipped"),
    ("6  ", "diff members \u2192 toAdd / toRemove sets"),
    ("7  ", "write in $batch \u2013 20 membership ops per call"),
    ("8  ", "sweep codes with zero failures \u2013 empty (delete only with -Cleanup)"),
    ("9  ", "report every count \u2192 any real failure throws \u2192 alert fires"),
], tcolor=PRIMARY_D, tsize=28, bsize=21)

card(M + cw + cgap, col_y, cw, col_h, "Safety: risky default \u2192 this repo",
     [
    ("DISM /ResetBase ", "\u2192 StartComponentCleanup only"),
    ("chkdsk /r scheduled ", "\u2192 read-only scan, /f behind -AllowChkdsk"),
    ("kill every service / msiexec ", "\u2192 age-gated (>30 min), top-5 report"),
    ("netsh winsock reset ", "\u2192 DNS cache + guidance only"),
    ("delete WinSxS\\pending.xml ", "\u2192 never touched"),
    ("always exit 0 ", "\u2192 Verified / Unknown \u2192 0, Failed \u2192 1"),
    ("-DryRun ", "previews every create/add/remove before it happens"),
    ("ownership marker ", "protects groups made by humans"),
], tcolor=ORANGE, fill=ORANGE_BG, tsize=28, bsize=21)

card(M + 2 * (cw + cgap), col_y, cw, col_h, "Contract & coverage", [
    ("detection exit 0 ", "= compliant, exit 1 = run the fix"),
    ("remediation exit 0 ", "= Verified / Unknown, exit 1 = Failed"),
    ("report line ", "= State|Detail per device in Intune"),
    ("20+ patch codes ", "+ 35+ app codes (MSI 1601\u20131642, 0x87D0\u2026)"),
    ("unknown code ", "= compliant \u2013 never loops on what we can't verify"),
    ("logs ", "= %TEMP%\\IntuneDetection.log / IntuneRemediation.log"),
    ("runbook failure ", "\u2192 job Failed \u2192 Azure Monitor alert"),
    ("offline proof ", "= Tests\\Run-Tests.ps1, 40/40 assertions"),
], tcolor=GREEN, fill=GREEN_BG, tsize=28, bsize=21)

# ============================================================ 3 - stats
chip_y = col_y + col_h + 0.45
chips = [("20+", "patch codes covered"), ("35+", "app codes covered"),
         ("40/40", "offline tests pass"), ("0", "secrets stored"),
         ("20", "ops per $batch call"), ("6", "retries on 429/5xx"),
         ("1000", "rows per Graph page"), ("daily", "06:00 schedule")]
cg = 0.35
chw = (CW - (len(chips) - 1) * cg) / len(chips)
for i, (big, small) in enumerate(chips):
    x = M + i * (chw + cg)
    c = box(x, chip_y, chw, 1.6, CARD_BG, CARD_BD)
    tf = c.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, True, big, 40, PRIMARY_D, True, PP_ALIGN.CENTER, space_after=2)
    _para(tf, False, small, 18, MUTED, False, PP_ALIGN.CENTER, space_after=0)

# ============================================================ 4 - rollout
roll_y = chip_y + 1.6 + 0.45
roll_w = CW * 0.60
r = box(M, roll_y, roll_w, 1.4, GREEN_BG, CARD_BD)
tf = r.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.3)
tf.margin_right = Inches(0.3)
_para(tf, True, "Rollout: offline tests \u2192 -DryRun against prod \u2192 pilot one "
      "code \u2192 full daily sync \u2192 -CleanupOrphanedGroups last",
      23, GREEN, True, PP_ALIGN.CENTER, space_after=4)
_para(tf, False, "Alerting armed before Phase 3; two clean weeks of job runs "
      "before cleanup is enabled", 20, TEXT, False, PP_ALIGN.CENTER,
      space_after=0)

lw = CW - roll_w - 0.6
r = box(M + roll_w + 0.6, roll_y, lw, 1.4, ORANGE_BG, CARD_BD)
tf = r.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.3)
tf.margin_right = Inches(0.3)
_para(tf, True, "Licensing checkpoint: Intune Remediations require Windows "
      "Enterprise E3/E5, Education A3/A5 or VDA per user",
      23, ORANGE, True, PP_ALIGN.CENTER, space_after=4)
_para(tf, False, "Beta endpoints (windowsUpdateStates / deviceStatuses) are "
      "documented, no SLA \u2013 isolated behind one helper",
      20, TEXT, False, PP_ALIGN.CENTER, space_after=0)

# ============================================================ footer
foot_y = H - 1.35
box(0, foot_y, W, 1.35, NAVY, None, MSO_SHAPE.RECTANGLE)
text_at(M, foot_y + 0.30, CW * 0.55, 0.8,
        [("Intune Patch & App Remediation Pipeline  \u00b7  "
          "Azure Automation \u00b7 Microsoft Graph \u00b7 Entra groups \u00b7 "
          "Intune Remediations", 21, RGBColor(0x9F, 0xC4, 0xE8), False, 0)])
text_at(M + CW * 0.55, foot_y + 0.30, CW * 0.45, 0.8,
        [("Regenerate: python Tests\\Build-Poster.py  \u00b7  "
          "python Tests\\Build-Presentation.py", 21,
          RGBColor(0x9F, 0xC4, 0xE8), False, 0)],
        align=PP_ALIGN.RIGHT)

s.notes_slide.notes_text_frame.text = (
    "A1 landscape poster companion to the 21-slide deck. Regenerate with "
    "python Tests\\Build-Poster.py, then export PNG via PowerPoint "
    "(File > Export, or Tests\\Export-PosterPng.ps1).")

prs.core_properties.title = "Intune Patch & App Remediation Pipeline - Poster"
prs.core_properties.author = "IntuneRemediations"
prs.save(OUT)
print("Saved:", OUT)
print("Slide size: %.2f x %.2f in" % (prs.slide_width / 914400,
                                      prs.slide_height / 914400))
